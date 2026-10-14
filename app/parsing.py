"""
Turn an uploaded file (PDF, DOCX, TXT, MD) into plain text "pages".
Each page is (page_number, text). PDFs keep real page numbers so citations can say "page 3".
Markdown/TXT files may start with front matter (title, allowed_groups, ...) between --- lines.
"""
import io
from dataclasses import dataclass, field

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class ParseError(ValueError):
    """Raised when a file cannot be read. We never index empty content silently."""


@dataclass
class ParsedDocument:
    pages: list  # list of (page_number, text)
    front_matter: dict = field(default_factory=dict)


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def split_front_matter(text: str):
    """Return (metadata_dict, body_text) for text that starts with '---' front matter."""
    if not text.startswith("---"):
        return {}, text
    lines = text.splitlines()
    meta = {}
    for i, line in enumerate(lines[1:], start=1):
        if line.strip() == "---":
            return meta, "\n".join(lines[i + 1:])
        if ":" in line:
            key, value = line.split(":", 1)
            meta[key.strip().lower()] = value.strip()
    return {}, text  # no closing '---' -> treat everything as body


def parse_file(filename: str, data: bytes) -> ParsedDocument:
    ext = _extension(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError(f"Unsupported file type '{ext}'. Allowed: {sorted(SUPPORTED_EXTENSIONS)}")

    if ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise ParseError("File has a .pdf name but is not a PDF")
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(i + 1, (p.extract_text() or "").strip()) for i, p in enumerate(reader.pages)]
        doc = ParsedDocument(pages=[(n, t) for n, t in pages if t])

    elif ext == ".docx":
        if not data.startswith(b"PK"):
            raise ParseError("File has a .docx name but is not a Word document")
        import docx  # python-docx

        d = docx.Document(io.BytesIO(data))
        parts = []
        for para in d.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            parts.append(f"## {text}" if style.startswith("heading") else text)
        for table in d.tables:  # keep each table row together on one line
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        doc = ParsedDocument(pages=[(1, "\n\n".join(parts))])

    else:  # .txt / .md
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        meta, body = split_front_matter(text)
        doc = ParsedDocument(pages=[(1, body.strip())], front_matter=meta)

    if not any(t.strip() for _, t in doc.pages):
        raise ParseError("No readable text found (scanned PDF? OCR is not enabled in this version)")
    return doc